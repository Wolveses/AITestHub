# -*- coding: utf-8 -*-
"""Excel 表格 Loader（.xlsx / .xls）。

高性能方案：优先使用 python-calamine（Rust 内核，毫秒级读取，超低内存占用），
若遇到不可解析异常或依赖缺失，平滑降级到 openpyxl。

特性：
1. 采用标准 Markdown 表格语法格式化输出（表头 + 分隔线 + 纯数据行），杜绝单行内部重复表头；
2. 智能复合多级表头识别与向上层级展开合并，消除跨列指标列丢失问题；
3. 清洗单元格换行，格式化日期时间与数值；
4. 保留 Sheet 结构与表头信息，meta 带 type="table" 供 TableChunker 感知。
"""

from __future__ import annotations

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Any

from insight_aitest.platform.services.kb.loader import register_loader
from insight_aitest.platform.services.kb.loader.base import DocumentLoadError, DocumentLoader
from insight_aitest.platform.services.kb.models import ParsedDocument

logger = logging.getLogger(__name__)


def _format_cell(val: Any) -> str:
    """归一化单元格内容为适合表格呈现的单行文本。"""
    if val is None:
        return ""
    if isinstance(val, bool):
        return "是" if val else "否"
    if isinstance(val, (datetime, date)):
        return val.isoformat()
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val))
        return f"{val:.4f}".rstrip("0").rstrip(".")
    s = str(val).strip()
    if not s:
        return ""
    # 将内部换行符转换为空格，转义管道符，保持单行 Markdown 表格列结构
    cleaned = " ".join(s.splitlines())
    return cleaned.replace("|", "\\|")


def _detect_and_build_headers(rows: list[list[str]]) -> tuple[int, list[str]]:
    """智能表头识别与多级复合表头融合。

    返回: (header_rows_count, final_headers)
    """
    if not rows:
        return 0, []

    max_cols = max(len(r) for r in rows)

    # 1. 探测表头行数（前 1~3 行是否存在复合子表头）
    header_rows_count = 1
    if len(rows) > 1 and any(not c for c in rows[0]):
        sub_words = {"平均值", "最大值", "最小值", "描述", "单位", "备注", "第一次", "第二次", "第三次"}
        for idx in range(1, min(4, len(rows))):
            r = rows[idx]
            if any(c in sub_words for c in r if c):
                header_rows_count = idx + 1
            else:
                break

    header_rows = rows[:header_rows_count]

    # 2. 补齐列数
    norm_headers: list[list[str]] = []
    for hr in header_rows:
        row_str = list(hr)
        if len(row_str) < max_cols:
            row_str.extend([""] * (max_cols - len(row_str)))
        norm_headers.append(row_str)

    # 3. 父级向右展开填充（合并单元格逻辑）
    for l_idx in range(len(norm_headers) - 1):
        cur = ""
        for c_idx in range(max_cols):
            val = norm_headers[l_idx][c_idx]
            if val:
                cur = val
            elif cur and norm_headers[l_idx + 1][c_idx]:
                norm_headers[l_idx][c_idx] = cur

    # 4. 上下级表头拼接
    final_headers: list[str] = []
    for c_idx in range(max_cols):
        parts: list[str] = []
        for r in norm_headers:
            val = r[c_idx]
            if val and val not in parts:
                parts.append(val)
        final_headers.append("-".join(parts) if parts else f"列{c_idx+1}")

    return header_rows_count, final_headers


@register_loader(".xlsx")
@register_loader(".xls")
class XlsxLoader(DocumentLoader):
    """Excel 表格 Loader，支持 .xlsx 与 .xls。"""

    def load(self, path: Path) -> ParsedDocument:
        # 大文件防御上限放宽到 100MB（Calamine 处理 100MB Excel 仅需 1~2s）
        file_size = path.stat().st_size
        if file_size > 100 * 1024 * 1024:
            raise DocumentLoadError(
                f"Excel 文件过大 ({file_size // (1024 * 1024)}MB > 100MB)，超出系统处理限制。"
            )

        try:
            return self._load_via_calamine(path)
        except Exception as e:
            logger.warning("calamine 加载 Excel 失败，尝试降级 openpyxl: %s", e)
            try:
                return self._load_via_openpyxl(path)
            except Exception as e2:
                raise DocumentLoadError(f"解析 Excel 失败: {e2}") from e2

    def _load_via_calamine(self, path: Path) -> ParsedDocument:
        from python_calamine import CalamineWorkbook

        wb = CalamineWorkbook.from_path(str(path))
        parts: list[str] = []
        max_rows_per_sheet = 5000
        total_char_limit = 3000000

        for sheet_name in wb.sheet_names:
            sheet = wb.get_sheet_by_name(sheet_name)
            raw_rows = sheet.to_python()
            if not raw_rows:
                continue

            # 转换为清洗后的二维文本列表
            norm_rows: list[list[str]] = []
            max_cols = 0
            for r in raw_rows:
                row_str = [_format_cell(c) for c in r]
                while row_str and not row_str[-1]:
                    row_str.pop()
                if row_str:
                    max_cols = max(max_cols, len(row_str))
                norm_rows.append(row_str)

            if not norm_rows or max_cols == 0:
                continue

            for r in norm_rows:
                if len(r) < max_cols:
                    r.extend([""] * (max_cols - len(r)))

            # 表头识别与多级展开
            header_count, final_headers = _detect_and_build_headers(norm_rows)
            parts.append(f"# {sheet_name}")

            # 输出 Markdown 表头行与分隔行
            header_cells = [h if h else f"列{i+1}" for i, h in enumerate(final_headers)]
            parts.append("| " + " | ".join(header_cells) + " |")
            parts.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

            # 输出纯数据行（不再每行重复添加表头前缀）
            data_rows = norm_rows[header_count:]
            row_count = 0
            for r in data_rows:
                if not any(r):
                    continue
                row_count += 1
                if row_count > max_rows_per_sheet:
                    parts.append(f"... (已截断超过 {max_rows_per_sheet} 行的数据)")
                    break

                parts.append("| " + " | ".join(r) + " |")

                if sum(len(p) for p in parts) > total_char_limit:
                    parts.append("... (已达到最大字符提取上限)")
                    break

        content = "\n".join(parts)
        return ParsedDocument(
            filename=path.name,
            content=content,
            meta={"format": "xlsx", "type": "table"},
        )

    def _load_via_openpyxl(self, path: Path) -> ParsedDocument:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), data_only=True)
        parts: list[str] = []
        max_rows_per_sheet = 5000
        total_char_limit = 3000000

        try:
            for sheet in wb.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                norm_rows: list[list[str]] = []
                max_cols = 0
                for r in rows:
                    row_str = [_format_cell(c) for c in r]
                    while row_str and not row_str[-1]:
                        row_str.pop()
                    if row_str:
                        max_cols = max(max_cols, len(row_str))
                    norm_rows.append(row_str)

                if not norm_rows or max_cols == 0:
                    continue

                for r in norm_rows:
                    if len(r) < max_cols:
                        r.extend([""] * (max_cols - len(r)))

                header_count, final_headers = _detect_and_build_headers(norm_rows)
                parts.append(f"# {sheet.title}")

                header_cells = [h if h else f"列{i+1}" for i, h in enumerate(final_headers)]
                parts.append("| " + " | ".join(header_cells) + " |")
                parts.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

                data_rows = norm_rows[header_count:]
                row_count = 0
                for r in data_rows:
                    if not any(r):
                        continue
                    row_count += 1
                    if row_count > max_rows_per_sheet:
                        parts.append(f"... (已截断超过 {max_rows_per_sheet} 行的数据)")
                        break

                    parts.append("| " + " | ".join(r) + " |")

                    if sum(len(p) for p in parts) > total_char_limit:
                        parts.append("... (已达到最大字符提取上限)")
                        break
        finally:
            wb.close()

        content = "\n".join(parts)
        return ParsedDocument(
            filename=path.name,
            content=content,
            meta={"format": "xlsx", "type": "table"},
        )
