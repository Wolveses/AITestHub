# -*- coding: utf-8 -*-
"""用例导出 API（支持 JSON 与 TAPD 规范 Excel 导出）。"""

from __future__ import annotations

import io
import re
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel

from insight_aitest.modules.testcase.backend.deps import get_tc_db
from insight_aitest.modules.testcase.backend.persistence.database import TestCaseDatabase
from insight_aitest.modules.testcase.backend.persistence.models import TestCase
from insight_aitest.modules.testcase.backend.routes.testcases import _out

router = APIRouter(prefix="/testcases", tags=["testcase"])


class ExportRequest(BaseModel):
    ids: list[int]


class ExportExcelRequest(BaseModel):
    case_ids: list[int] | None = None
    batch_id: str | None = None
    task_id: int | None = None


TAPD_HEADERS = [
    "用例目录",
    "用例名称",
    "需求ID",
    "前置条件",
    "用例步骤",
    "预期结果",
    "用例类型",
    "用例状态",
    "用例等级",
    "创建人",
]

TAPD_DESC_ROW = [
    "[字段填写说明]“用例目录”请填写完整路径，用“-”分隔。如果目录为空，默认导入为“未规划目录”中；如果用例目录不存在，请在预览页面选择是否要自动创建目录。",
    "“用例名称”为必填项。",
    "“需求ID”请填写需求ID，多个需求ID以英文;号隔开。需求必须是本项目下的需求。",
    "“前置条件”请填写合法文本。",
    "“用例步骤”请填写合法文本。",
    "“预期结果”请填写合法文本。",
    "“用例类型”请填写：功能测试、性能测试、安全性测试、其他、冒烟测试。",
    "“用例状态”请填写：正常、待更新、已废弃。",
    "“用例等级”请填写：高、中、低。",
    '支持多个人员传入,使用\';\'隔开 如: "xxx;xxx"',
]

TYPE_MAP = {
    "functional": "功能测试",
    "performance": "性能测试",
    "security": "安全性测试",
    "smoke": "冒烟测试",
    "api": "功能测试",
    "ui": "功能测试",
}

STATUS_MAP = {
    "deprecated": "已废弃",
    "draft": "正常",
    "reviewed": "正常",
    "ready": "正常",
}

PRIORITY_MAP = {
    "p0": "高",
    "p1": "高",
    "p2": "中",
    "p3": "低",
}


def _format_steps(content: dict) -> str:
    if not isinstance(content, dict):
        return ""
    steps = content.get("steps")
    if not steps:
        return ""
    if isinstance(steps, str):
        lines = [l.strip() for l in steps.split("\n") if l.strip()]
        if len(lines) <= 1:
            return steps.strip()
        numbered = []
        for idx, l in enumerate(lines, 1):
            if re.match(r"^\d+[、\.．]", l):
                numbered.append(l)
            else:
                numbered.append(f"{idx}、{l}")
        return "\n".join(numbered)

    if isinstance(steps, list):
        numbered = []
        for idx, s in enumerate(steps, 1):
            if isinstance(s, dict):
                act = s.get("action") or s.get("desc") or s.get("step") or s.get("name") or str(s)
                numbered.append(f"{idx}、{act}")
            elif isinstance(s, str):
                clean_s = s.strip()
                if re.match(r"^\d+[、\.．]", clean_s):
                    numbered.append(clean_s)
                else:
                    numbered.append(f"{idx}、{clean_s}")
        return "\n".join(numbered)
    return str(steps)


def _format_expected(content: dict) -> str:
    if not isinstance(content, dict):
        return ""
    exp = content.get("expected")
    if exp:
        if isinstance(exp, str):
            lines = [l.strip() for l in exp.split("\n") if l.strip()]
            if len(lines) <= 1:
                return exp.strip()
            numbered = []
            for idx, l in enumerate(lines, 1):
                clean_l = re.sub(r"^\d+[、\.．]", "", l).strip()
                numbered.append(f"{idx}、{clean_l}")
            return "\n".join(numbered)
        if isinstance(exp, list):
            return "\n".join(f"{idx}、{str(e).strip()}" for idx, e in enumerate(exp, 1))

    steps = content.get("steps")
    if isinstance(steps, list):
        step_exps = []
        for idx, s in enumerate(steps, 1):
            if isinstance(s, dict) and s.get("expected"):
                step_exps.append(f"{idx}、{s['expected']}")
        if step_exps:
            return "\n".join(step_exps)

    return ""


def generate_tapd_excel(cases: list[TestCase], db: TestCaseDatabase) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试用例"

    ws.append(TAPD_HEADERS)

    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    header_font = Font(bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 45
    ws.column_dimensions["F"].width = 45
    ws.column_dimensions["G"].width = 15
    ws.column_dimensions["H"].width = 15
    ws.column_dimensions["I"].width = 15
    ws.column_dimensions["J"].width = 15

    folders = db.list_folders()
    folder_map = {f.id: f for f in folders}

    def _get_folder_path(fid: int | None) -> str:
        if not fid or fid not in folder_map:
            return ""
        parts = []
        curr = folder_map[fid]
        visited = set()
        while curr and curr.id not in visited:
            visited.add(curr.id)
            parts.append(curr.name)
            curr = folder_map.get(curr.parent_id) if curr.parent_id else None
        return "-".join(reversed(parts))

    for c in cases:
        type_val = c.type.value if hasattr(c.type, "value") else str(c.type)
        status_val = c.status.value if hasattr(c.status, "value") else str(c.status)
        priority_val = c.priority.value if hasattr(c.priority, "value") else str(c.priority)

        row_data = [
            _get_folder_path(c.folder_id),
            c.title,
            "",
            c.preconditions or "",
            _format_steps(c.content),
            _format_expected(c.content),
            TYPE_MAP.get(type_val, "功能测试"),
            STATUS_MAP.get(status_val, "正常"),
            PRIORITY_MAP.get(priority_val, "中"),
            "",
        ]
        ws.append(row_data)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _export_dict(case) -> dict:
    """导出完整用例结构（含 content/preconditions/status 等全字段）。"""
    return _out(case).model_dump()


@router.get("/{case_id}/export")
async def export_one(case_id: int, db: TestCaseDatabase = Depends(get_tc_db)) -> dict:
    case = db.get_case(case_id)
    if not case:
        raise HTTPException(404, "用例不存在")
    return _export_dict(case)


@router.post("/export")
async def export_batch(
    body: ExportRequest, db: TestCaseDatabase = Depends(get_tc_db)
) -> list[dict]:
    out = []
    for cid in body.ids:
        case = db.get_case(cid)
        if case:
            out.append(_export_dict(case))
    return out


@router.post("/export-excel")
async def export_excel(
    body: ExportExcelRequest, db: TestCaseDatabase = Depends(get_tc_db)
):
    cases: list[TestCase] = []
    if body.case_ids:
        for cid in body.case_ids:
            c = db.get_case(cid)
            if c:
                cases.append(c)
    elif body.batch_id:
        cases = db.list_cases_by_batch(body.batch_id)
    elif body.task_id:
        cases = db.list_cases(task_id=body.task_id)

    if not cases:
        raise HTTPException(400, "没有可导出的用例")

    from datetime import datetime
    from urllib.parse import quote

    now_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"测试用例_{now_str}.xlsx"
    quoted_name = quote(filename)

    data = generate_tapd_excel(cases, db)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="testcases_{now_str}.xlsx"; filename*=UTF-8\'\'{quoted_name}',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
