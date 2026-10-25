# -*- coding: utf-8 -*-
"""导出 API 测试。"""
from fastapi import FastAPI
from fastapi.testclient import TestClient


def _setup_app(tmp_path, monkeypatch):
    monkeypatch.setenv("INSIGHT_EYE_AI_LLM_API_KEY", "sk-test")
    monkeypatch.setenv("INSIGHT_EYE_AI_EMBED_DIM", "4")
    import insight_aitest.modules.testcase.backend.deps as tc_deps
    tc_deps._tc_db = None
    from insight_aitest.modules.testcase.backend.persistence.database import TestCaseDatabase
    tc_deps._tc_db = TestCaseDatabase(str(tmp_path / "testcase.db"))
    from insight_aitest.modules.testcase.backend.routes import router as tc_router
    app = FastAPI()
    app.include_router(tc_router, prefix="/api/modules/testcase")
    return TestClient(app)


def test_export_one(tmp_path, monkeypatch):
    c = _setup_app(tmp_path, monkeypatch)
    cid = c.post("/api/modules/testcase/testcases", json={
        "title": "导出测试", "type": "functional",
        "content": {"steps": [{"no": 1, "action": "x", "data": ""}], "expected": "y"},
        "preconditions": "已开页"}).json()["id"]
    r = c.get(f"/api/modules/testcase/testcases/{cid}/export")
    assert r.status_code == 200
    body = r.json()
    assert body["content"]["expected"] == "y"
    assert body["preconditions"] == "已开页"
    assert body["type"] == "functional"


def test_export_batch(tmp_path, monkeypatch):
    c = _setup_app(tmp_path, monkeypatch)
    id1 = c.post("/api/modules/testcase/testcases", json={"title": "a"}).json()["id"]
    id2 = c.post("/api/modules/testcase/testcases", json={"title": "b"}).json()["id"]
    r = c.post("/api/modules/testcase/testcases/export", json={"ids": [id1, id2, 999]})
    assert r.status_code == 200
    body = r.json()
    assert len(body) == 2  # 999 不存在，跳过


def test_export_not_found(tmp_path, monkeypatch):
    c = _setup_app(tmp_path, monkeypatch)
    assert c.get("/api/modules/testcase/testcases/999/export").status_code == 404


def test_export_excel(tmp_path, monkeypatch):
    import openpyxl
    c = _setup_app(tmp_path, monkeypatch)
    # 创建目录
    fid = c.post("/api/modules/testcase/testcases/folders", json={"name": "搜索模块"}).json()["id"]
    # 创建用例
    id1 = c.post("/api/modules/testcase/testcases", json={
        "title": "测试搜索功能",
        "type": "functional",
        "priority": "p0",
        "preconditions": "用户已登录",
        "folder_id": fid,
        "content": {
            "steps": ["输入关键词", "点击搜索按钮"],
            "expected": "显示匹配结果",
        },
    }).json()["id"]

    r = c.post("/api/modules/testcase/testcases/export-excel", json={"case_ids": [id1]})
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "filename*=UTF-8''" in r.headers["content-disposition"]

    # 验证 openpyxl 可以正常解析且符合格式（仅保留标题行和用例数据，无注释行）
    import io
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.title == "测试用例"
    # Row 1 表头
    headers = [ws.cell(row=1, column=i).value for i in range(1, 11)]
    assert headers[0] == "用例目录"
    assert headers[1] == "用例名称"
    assert headers[6] == "用例类型"
    assert headers[8] == "用例等级"
    # Row 2 直接为用例数据（已删除注释说明行）
    assert ws.cell(row=2, column=1).value == "搜索模块"
    assert ws.cell(row=2, column=2).value == "测试搜索功能"
    assert ws.cell(row=2, column=4).value == "用户已登录"
    assert "1、输入关键词" in str(ws.cell(row=2, column=5).value)
    assert "显示匹配结果" in str(ws.cell(row=2, column=6).value)
    assert ws.cell(row=2, column=7).value == "功能测试"
    assert ws.cell(row=2, column=8).value == "正常"
    assert ws.cell(row=2, column=9).value == "高"
